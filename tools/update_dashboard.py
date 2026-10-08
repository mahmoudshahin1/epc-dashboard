#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
EPC Production Dashboard — Builder
==================================
Reads the ERP production export (0040_qa.xls / .xlsx) and builds:

  dist/                 -> multi-file website (for Netlify / any server)
      index.html
      css/style.css
      js/app.js
      js/data.js        <- generated (data + config + fingerprint)
      js/vendor/xlsx.full.min.js
      assets/logo.png
      standalone.html   <- single self-contained file (WhatsApp / offline / USB)

Project layout expected (this file lives in  <project>/tools/):

  <project>/
    index.html  css/style.css  js/app.js  js/vendor/...  assets/logo.png
    plan.json   update_dashboard.bat  publish.bat  README*.txt
    tools/update_dashboard.py

Usage:
    python tools/update_dashboard.py                     # auto-find export
    python tools/update_dashboard.py C:\\path\\0040_qa.xls
"""
import sys
import json
import re
import argparse
import calendar
import datetime
import hashlib
import base64
import shutil
from collections import Counter
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
PLAN_PATH = PROJECT_ROOT / "plan.json"
DIST = PROJECT_ROOT / "dist"

AR_MONTHS = {1: "يناير", 2: "فبراير", 3: "مارس", 4: "أبريل", 5: "مايو", 6: "يونيو",
             7: "يوليو", 8: "أغسطس", 9: "سبتمبر", 10: "أكتوبر", 11: "نوفمبر", 12: "ديسمبر"}

DEFAULT_PLAN = {
    "plan_total_tons": 4442,
    "plan_AL_tons": 3139,
    "plan_CU_tons": 1303,
    "for_month": "",
    "_comment": "Edit these numbers every month. for_month is optional: 'YYYY-MM' to warn if data month differs."
}

REQUIRED_COLS = {"wo", "export", "cond_kg", "screen_kgm", "armour_kgm", "prodfamily",
                 "customer", "packno", "packqty", "prdate", "maccode", "conductor", "qcstatus"}


def die(msg):
    print("ERROR: " + msg)
    sys.exit(1)


def find_export(cli_path):
    if cli_path:
        p = Path(cli_path)
        if not p.exists():
            die(f"file not found: {p}")
        return p
    names = ["0040_qa.xls", "0040_qa.xlsx", "0040_qa.xlsm"]
    searched = []
    for base in (PROJECT_ROOT, Path.home() / "Desktop", Path.home() / "Desktop" / "5-10"):
        for n in names:
            cand = base / n
            searched.append(str(cand))
            if cand.exists():
                return cand
    for pat in ("*.xls", "*.xlsx", "*.xlsm"):
        for p in sorted(PROJECT_ROOT.glob(pat)):
            if p.name.startswith("~$"):
                continue
            return p
    die("could not find the ERP export file (0040_qa.xls/.xlsx).\n"
        "Put it in the project folder, on the Desktop, or pass the full path:\n"
        "    python tools/update_dashboard.py \"C:\\path\\to\\0040_qa.xls\"\n"
        "Searched:\n  " + "\n  ".join(searched))


def read_rows(path):
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            die("reading .xlsx needs openpyxl. Install with:  pip install openpyxl")
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        ws = wb["0040_qa"] if "0040_qa" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows, None
    if ext == ".xls":
        try:
            import xlrd
        except ImportError:
            die("reading old .xls needs xlrd. Install with:  pip install xlrd")
        import xlrd as _xlrd
        book = _xlrd.open_workbook(str(path))
        names = book.sheet_names()
        sh = book.sheet_by_name("0040_qa") if "0040_qa" in names else book.sheet_by_index(0)
        rows = [sh.row_values(i) for i in range(sh.nrows)]
        return rows, book.datemode
    die(f"unsupported file type '{ext}' — need .xls or .xlsx")


def find_header(rows):
    for i, row in enumerate(rows[:10]):
        cells = {str(v).strip().lower() for v in row if v is not None}
        if REQUIRED_COLS.issubset(cells):
            return i, {str(v).strip().lower(): j for j, v in enumerate(row) if v is not None}
    die("could not find the header row (expected columns like wo/packno/prdate/cond_kg...). "
        "Is this the right ERP export file?")


def to_date(v, datemode):
    if v is None or v == "":
        return None
    if hasattr(v, "year") and hasattr(v, "month") and hasattr(v, "day"):
        return v
    if isinstance(v, (int, float)) and v > 20000:
        if datemode is not None:
            import xlrd
            try:
                return xlrd.xldate.xldate_as_datetime(float(v), datemode)
            except Exception:
                return None
        return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(v))
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
            try:
                return datetime.datetime.strptime(s[:19], fmt)
            except ValueError:
                pass
    return None


def fnum(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def sval(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def extract(rows, hdr_idx, datemode):
    def g(r, name):
        j = hdr_idx.get(name)
        return r[j] if j is not None and j < len(r) else None

    data = []
    skipped = 0
    for r in rows:
        d = to_date(g(r, "prdate"), datemode)
        pk = sval(g(r, "packno"))
        if d is None or pk == "":
            skipped += 1
            continue
        q = fnum(g(r, "packqty"))
        ct = q * fnum(g(r, "cond_kg")) / 1000.0
        st = q * fnum(g(r, "screen_kgm")) / 1000.0
        at = q * fnum(g(r, "armour_kgm")) / 1000.0
        data.append({
            "d": d.day, "pk": pk, "wo": sval(g(r, "wo")),
            "cu": sval(g(r, "customer")) or "?",
            "fam": sval(g(r, "prodfamily")) or "?",
            "con": sval(g(r, "conductor")) or "?",
            "exp": "Y" if sval(g(r, "export")).upper() == "Y" else "N",
            "q": round(q), "ct": round(ct, 4), "st": round(st, 4), "at": round(at, 4),
            "tt": round(ct + st + at, 4),
            "mac": sval(g(r, "maccode")) or "?",
            "pt": sval(g(r, "packtype")) or "?",
            "qc": sval(g(r, "qcstatus")) or "P",
            "vs": (sval(g(r, "volt")) + " " + sval(g(r, "size"))).strip(),
            "_ym": (d.year, d.month),
        })
    if not data:
        die("no production rows found in the export file (no valid prdate/packno).")
    return data, skipped


def load_plan():
    if not PLAN_PATH.exists():
        PLAN_PATH.write_text(json.dumps(DEFAULT_PLAN, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[plan] Created {PLAN_PATH.name} with default values — edit it with Notepad to set your monthly plan.")
        return dict(DEFAULT_PLAN)
    try:
        return json.loads(PLAN_PATH.read_text(encoding="utf-8"))
    except Exception as e:
        die(f"plan.json is not valid JSON ({e}). Fix it or delete it to regenerate defaults.")


def safe_js(s):
    return re.sub(r"</script", r"<\\/script", s, flags=re.I)


def build_standalone():
    html = (DIST / "index.html").read_text(encoding="utf-8")
    css = (DIST / "css" / "style.css").read_text(encoding="utf-8")
    app = (DIST / "js" / "app.js").read_text(encoding="utf-8")
    vend = (DIST / "js" / "vendor" / "xlsx.full.min.js").read_text(encoding="utf-8")
    datajs = (DIST / "js" / "data.js").read_text(encoding="utf-8")
    logo_b64 = base64.b64encode((PROJECT_ROOT / "assets" / "logo.png").read_bytes()).decode()
    html = html.replace('<link rel="stylesheet" href="css/style.css">', "<style>\n" + css + "</style>")
    html = html.replace('<script src="js/vendor/xlsx.full.min.js"></script>', "<script>" + safe_js(vend) + "</script>")
    html = html.replace('<script src="js/data.js"></script>', "<script>" + safe_js(datajs) + "</script>")
    html = html.replace('<script src="js/app.js"></script>', "<script>" + safe_js(app) + "</script>")
    html = html.replace('src="assets/logo.png"', 'src="data:image/png;base64,' + logo_b64 + '"')
    for needle in ('href="css/', 'src="js/', 'src="assets/'):
        if needle in html:
            die(f"standalone build still references external file: {needle}")
    return html


def main():
    ap = argparse.ArgumentParser(description="Build the EPC production dashboard from the ERP export.")
    ap.add_argument("export", nargs="?", help="path to 0040_qa.xls / .xlsx (default: auto-find)")
    args = ap.parse_args()

    for need in ("index.html", "css/style.css", "js/app.js", "js/vendor/xlsx.full.min.js", "assets/logo.png"):
        if not (PROJECT_ROOT / need).exists():
            die(f"project file missing: {need} — keep the project structure intact.")

    src = find_export(args.export)
    rows, datemode = read_rows(src)
    hdr_row, hdr_idx = find_header(rows)
    data, skipped = extract(rows[hdr_row + 1:], hdr_idx, datemode)

    ym = Counter(r["_ym"] for r in data).most_common(1)[0][0]
    year, month = ym
    mdays = calendar.monthrange(year, month)[1]
    for r in data:
        del r["_ym"]

    plan = load_plan()
    plan_total = fnum(plan.get("plan_total_tons"))
    plan_al = fnum(plan.get("plan_AL_tons"))
    plan_cu = fnum(plan.get("plan_CU_tons"))
    for_month = str(plan.get("for_month", "")).strip()
    if for_month and for_month != f"{year:04d}-{month:02d}":
        print(f"[plan] WARNING: plan.json for_month='{for_month}' but data month is "
              f"{year:04d}-{month:02d} — make sure the plan matches this month!")

    raw_json = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    config = {
        "plan_total": plan_total, "plan_al": plan_al, "plan_cu": plan_cu,
        "month_ar": AR_MONTHS.get(month, str(month)), "month_num": month,
        "year": year, "month_days": mdays, "source": src.name,
        "generated": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "fingerprint": hashlib.sha256(raw_json.encode("utf-8")).hexdigest()[:10],
    }

    # ---- dist (multi-file site) ----
    if DIST.exists():
        shutil.rmtree(DIST)
    (DIST / "css").mkdir(parents=True)
    (DIST / "js" / "vendor").mkdir(parents=True)
    (DIST / "assets").mkdir(parents=True)
    shutil.copy(PROJECT_ROOT / "index.html", DIST / "index.html")
    shutil.copy(PROJECT_ROOT / "css" / "style.css", DIST / "css" / "style.css")
    shutil.copy(PROJECT_ROOT / "js" / "app.js", DIST / "js" / "app.js")
    shutil.copy(PROJECT_ROOT / "js" / "vendor" / "xlsx.full.min.js", DIST / "js" / "vendor" / "xlsx.full.min.js")
    shutil.copy(PROJECT_ROOT / "assets" / "logo.png", DIST / "assets" / "logo.png")
    datajs = ("window.DASH_DATA=" + raw_json.replace("</", "<\\/") + ";\n"
              "window.DASH_CONFIG=" + json.dumps(config, ensure_ascii=False).replace("</", "<\\/") + ";\n")
    (DIST / "js" / "data.js").write_text(datajs, encoding="utf-8")

    # ---- standalone single file ----
    (DIST / "standalone.html").write_text(build_standalone(), encoding="utf-8")

    # ---- console summary ----
    total = sum(r["tt"] for r in data)
    al = sum(r["ct"] for r in data if r["con"] == "A")
    cu = sum(r["ct"] for r in data if r["con"] == "C")
    scr = sum(r["st"] for r in data)
    arm = sum(r["at"] for r in data)
    km = sum(r["q"] for r in data) / 1000.0
    max_day = max(r["d"] for r in data)
    days_left = max(0, mdays - max_day)
    remaining = max(0.0, plan_total - total) if plan_total else 0.0
    need = remaining / days_left if (plan_total and days_left) else 0.0
    n_days = len({r["d"] for r in data})
    avg = total / n_days if n_days else 0.0

    print("=" * 62)
    print(" EPC PRODUCTION DASHBOARD - BUILT")
    print("=" * 62)
    print(f" Source    : {src}  ({len(data)} packs" + (f", {skipped} rows skipped" if skipped else "") + ")")
    print(f" Month     : {year}-{month:02d} ({AR_MONTHS.get(month, month)}) - data days 1..{max_day} of {mdays}")
    print(f" Produced  : {total:,.2f} t   (AL {al:,.2f} / CU {cu:,.2f} / Screen {scr:,.2f} / Armour {arm:,.2f})")
    print(f" Length    : {km:,.1f} km")
    if plan_total:
        print(f" Plan      : {plan_total:,.0f} t  ->  DONE {total / plan_total * 100:,.1f}%  ({total:,.1f} t)")
        print(f" Remaining : {remaining:,.1f} t in {days_left} days -> need {need:,.1f} t/day "
              f"(current avg {avg:,.1f} t/day) {'=> ON PACE' if need <= avg * 1.05 else '=> BEHIND PACE'}")
    else:
        print(" Plan      : NOT SET - edit plan.json to enable plan tracking")
    print(f" Site      : {DIST}            (upload this folder to Netlify)")
    print(f" Standalone: {DIST / 'standalone.html'}   (send via WhatsApp / open offline)")
    print(f" Fingerprint: {config['fingerprint']}")
    print("=" * 62)


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    main()
